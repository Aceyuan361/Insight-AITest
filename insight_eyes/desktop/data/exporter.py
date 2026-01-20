"""
数据导出功能
支持导出为CSV、JSON、Excel格式，并生成性能测试报告
"""
import csv
import json
import os
from datetime import datetime
from typing import List, Dict, Optional
from pathlib import Path

from logzero import logger

from .database import DatabaseManager
from .repository import MetricsRepository, AlertRepository, SessionRepository


# 定义安全的导出目录
DEFAULT_EXPORT_DIR = os.path.abspath(os.path.expanduser('~/Documents/InsightEye'))


def validate_filepath(filepath: str, allowed_dir: str = None) -> str:
    """
    验证文件路径是否在允许的目录内，防止路径遍历攻击

    Args:
        filepath: 要验证的文件路径
        allowed_dir: 允许的目录，如果为None则使用默认导出目录

    Returns:
        规范化后的绝对路径

    Raises:
        ValueError: 如果路径不在允许的目录内
    """
    if allowed_dir is None:
        allowed_dir = DEFAULT_EXPORT_DIR

    # 规范化路径
    abs_filepath = os.path.abspath(filepath)
    abs_allowed_dir = os.path.abspath(allowed_dir)

    # 检查路径是否在允许的目录内
    if not abs_filepath.startswith(abs_allowed_dir):
        raise ValueError(
            f"Invalid file path: {filepath}. "
            f"Export path must be within {abs_allowed_dir}"
        )

    return abs_filepath


class DataExporter:
    """数据导出器 - 支持多种格式的数据导出"""

    def __init__(self, db: DatabaseManager):
        """
        初始化数据导出器

        Args:
            db: 数据库管理器实例
        """
        self.db = db
        self.metrics_repo = MetricsRepository(db)
        self.alert_repo = AlertRepository(db)
        self.session_repo = SessionRepository(db)

    def export_to_csv(self, session_id: int, filepath: str,
                      include_alerts: bool = True) -> bool:
        """
        导出为CSV格式

        Args:
            session_id: 会话ID
            filepath: 输出文件路径
            include_alerts: 是否包含告警数据

        Returns:
            是否成功
        """
        try:
            # 验证文件路径，防止路径遍历攻击
            validated_path = validate_filepath(filepath)

            # 确保目录存在
            os.makedirs(os.path.dirname(validated_path), exist_ok=True)

            # 获取会话信息
            session = self.db.get_session(session_id)
            if not session:
                logger.warning(f"会话 {session_id} 不存在")
                return False

            # 导出性能指标
            metrics = self.db.get_metrics(session_id)

            with open(validated_path, 'w', newline='', encoding='utf-8-sig') as f:
                if metrics:
                    # 处理指标数据：将时间戳转换为易读格式
                    formatted_metrics = []
                    for metric in metrics:
                        formatted_metric = {}
                        # 将时间戳放在第一列并格式化
                        if 'timestamp' in metric:
                            try:
                                ts = metric['timestamp']
                                if isinstance(ts, str):
                                    dt = datetime.fromisoformat(ts)
                                else:
                                    dt = datetime.fromtimestamp(ts)
                                formatted_metric['时间'] = dt.strftime('%Y-%m-%d %H:%M:%S.%f')[:-3]  # 毫秒精度
                            except:
                                formatted_metric['时间'] = str(metric['timestamp'])

                        # 添加其他字段（中文名称映射）
                        field_mapping = {
                            'id': 'ID',
                            'session_id': '会话ID',
                            'cpu_app': 'CPU使用率(%)',
                            'cpu_system': '系统CPU(%)',
                            'memory_app_private': '私有内存(MB)',
                            'memory_pss': 'PSS内存(MB)',
                            'memory_vss': 'VSS内存(MB)',
                            'fps': '帧率(FPS)',
                            'fps_jank_count': '卡顿次数',
                            'network_up_speed': '上行速率(KB/s)',
                            'network_down_speed': '下行速率(KB/s)',
                            'battery_level': '电池电量(%)',
                            'battery_temp': '电池温度(℃)',
                            'device_temp': '设备温度(℃)',
                            'network_type': '网络类型'
                        }

                        for key, value in metric.items():
                            if key == 'timestamp':
                                continue  # 已处理
                            if key in field_mapping:
                                formatted_metric[field_mapping[key]] = value

                        formatted_metrics.append(formatted_metric)

                    # 写入指标数据
                    if formatted_metrics:
                        writer = csv.DictWriter(f, fieldnames=list(formatted_metrics[0].keys()))
                        writer.writeheader()
                        writer.writerows(formatted_metrics)

            # 如果需要，导出告警数据到单独的文件
            if include_alerts:
                alerts_path = validated_path.replace('.csv', '_alerts.csv')
                alerts = self.db.get_alerts(session_id=session_id)

                with open(alerts_path, 'w', newline='', encoding='utf-8-sig') as f:
                    if alerts:
                        # 格式化告警数据
                        formatted_alerts = []
                        for alert in alerts:
                            formatted_alert = {}
                            # 时间戳格式化
                            if 'timestamp' in alert:
                                try:
                                    ts = alert['timestamp']
                                    if isinstance(ts, str):
                                        dt = datetime.fromisoformat(ts)
                                    else:
                                        dt = datetime.fromtimestamp(ts)
                                    formatted_alert['时间'] = dt.strftime('%Y-%m-%d %H:%M:%S')
                                except:
                                    formatted_alert['时间'] = str(alert['timestamp'])

                            # 其他字段映射
                            field_mapping = {
                                'id': 'ID',
                                'session_id': '会话ID',
                                'metric_type': '指标类型',
                                'severity': '严重程度',
                                'message': '描述',
                                'threshold': '阈值',
                                'current_value': '当前值'
                            }

                            for key, value in alert.items():
                                if key == 'timestamp':
                                    continue
                                if key in field_mapping:
                                    formatted_alert[field_mapping[key]] = value

                            formatted_alerts.append(formatted_alert)

                        writer = csv.DictWriter(f, fieldnames=list(formatted_alerts[0].keys()))
                        writer.writeheader()
                        writer.writerows(formatted_alerts)

            logger.info(f"CSV导出成功: {validated_path}")
            return True

        except ValueError as e:
            logger.error(f"路径验证失败: {e}")
            return False
        except Exception as e:
            logger.error(f"CSV导出失败: {e}")
            return False

    def export_to_json(self, session_id: int, filepath: str,
                       include_alerts: bool = True,
                       include_stats: bool = True) -> bool:
        """
        导出为JSON格式

        Args:
            session_id: 会话ID
            filepath: 输出文件路径
            include_alerts: 是否包含告警数据
            include_stats: 是否包含统计数据

        Returns:
            是否成功
        """
        try:
            # 验证文件路径，防止路径遍历攻击
            validated_path = validate_filepath(filepath)

            # 确保目录存在
            os.makedirs(os.path.dirname(validated_path), exist_ok=True)

            # 获取会话信息
            session_overview = self.session_repo.get_session_overview(session_id)
            if not session_overview:
                logger.warning(f"会话 {session_id} 不存在")
                return False

            # 构建导出数据
            export_data = {
                'session': session_overview['session_info'],
                'metrics': self.db.get_metrics(session_id)
            }

            # 添加统计数据
            if include_stats:
                export_data['statistics'] = self.metrics_repo.get_statistics(session_id)

            # 添加告警数据
            if include_alerts:
                export_data['alerts'] = self.db.get_alerts(session_id=session_id)
                export_data['alert_summary'] = self.alert_repo.get_alert_summary(session_id)

            # 写入JSON文件
            with open(validated_path, 'w', encoding='utf-8') as f:
                json.dump(export_data, f, ensure_ascii=False, indent=2, default=str)

            logger.info(f"JSON导出成功: {validated_path}")
            return True

        except ValueError as e:
            logger.error(f"路径验证失败: {e}")
            return False
        except Exception as e:
            logger.error(f"JSON导出失败: {e}")
            return False

    def export_to_excel(self, session_id: int, filepath: str) -> bool:
        """
        导出为Excel格式（多工作表）

        Args:
            session_id: 会话ID
            filepath: 输出文件路径

        Returns:
            是否成功
        """
        try:
            # 导入openpyxl
            try:
                from openpyxl import Workbook
                from openpyxl.styles import Font, Alignment, PatternFill
                from openpyxl.utils import get_column_letter
            except ImportError:
                print("错误: 需要安装 openpyxl 库")
                print("请运行: pip install openpyxl")
                return False

            # 验证文件路径（防止路径遍历攻击）
            try:
                validated_path = validate_filepath(filepath, self.DEFAULT_EXPORT_DIR)
                filepath = validated_path
            except ValueError as e:
                print(f"路径验证失败: {e}")
                return False

            # 确保目录存在
            dir_path = os.path.dirname(filepath)
            if dir_path:  # 只有当目录路径非空时才创建
                os.makedirs(dir_path, exist_ok=True)

            # 获取会话信息
            session = self.db.get_session(session_id)
            if not session:
                print(f"会话 {session_id} 不存在")
                return False

            # 创建工作簿
            wb = Workbook()
            wb.remove(wb.active)  # 删除默认工作表

            # 1. 会话概览工作表
            self._create_session_overview_sheet(wb, session_id)

            # 2. 性能指标工作表
            self._create_metrics_sheet(wb, session_id)

            # 3. 统计数据工作表
            self._create_statistics_sheet(wb, session_id)

            # 4. 告警记录工作表
            self._create_alerts_sheet(wb, session_id)

            # 5. 趋势数据工作表
            self._create_trends_sheet(wb, session_id)

            # 保存文件
            wb.save(filepath)
            print(f"Excel导出成功: {filepath}")
            return True

        except Exception as e:
            print(f"Excel导出失败: {e}")
            return False

    def _create_session_overview_sheet(self, wb, session_id: int):
        """创建会话概览工作表"""
        ws = wb.create_sheet("会话概览")

        overview = self.session_repo.get_session_overview(session_id)
        if not overview:
            return

        # 标题样式
        title_font = Font(bold=True, size=14)
        header_font = Font(bold=True, size=11)
        header_fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")

        # 添加标题
        ws['A1'] = "性能监控会话概览"
        ws['A1'].font = title_font

        row = 3

        # 会话基本信息
        ws[f'A{row}'] = "基本信息"
        ws[f'A{row}'].font = header_font
        row += 1

        session_info = overview['session_info']
        ws[f'A{row}'] = "会话ID"
        ws[f'B{row}'] = session_info['id']
        row += 1

        ws[f'A{row}'] = "设备名称"
        ws[f'B{row}'] = session_info['device_name']
        row += 1

        ws[f'A{row}'] = "平台"
        ws[f'B{row}'] = session_info['platform']
        row += 1

        ws[f'A{row}'] = "应用包名"
        ws[f'B{row}'] = session_info['package_name']
        row += 1

        ws[f'A{row}'] = "开始时间"
        ws[f'B{row}'] = session_info['start_time']
        row += 1

        ws[f'A{row}'] = "结束时间"
        ws[f'B{row}'] = session_info.get('end_time', '进行中')
        row += 1

        ws[f'A{row}'] = "监控时长"
        ws[f'B{row}'] = session_info['duration_formatted']
        row += 1

        ws[f'A{row}'] = "采样间隔"
        ws[f'B{row}'] = f"{session_info['sample_interval']}ms"
        row += 1

        # 数据摘要
        row += 1
        ws[f'A{row}'] = "数据摘要"
        ws[f'A{row}'].font = header_font
        row += 1

        data_summary = overview['data_summary']
        ws[f'A{row}'] = "总样本数"
        ws[f'B{row}'] = data_summary['total_samples']
        row += 1

        ws[f'A{row}'] = "告警总数"
        ws[f'B{row}'] = data_summary['alert_count']
        row += 1

        ws[f'A{row}'] = "未解决告警"
        ws[f'B{row}'] = data_summary['unresolved_alerts']
        row += 1

        # 设置列宽
        ws.column_dimensions['A'].width = 20
        ws.column_dimensions['B'].width = 40

    def _create_metrics_sheet(self, wb, session_id: int):
        """创建性能指标工作表"""
        ws = wb.create_sheet("性能指标")

        metrics = self.db.get_metrics(session_id)
        if not metrics:
            return

        # 写入表头
        headers = list(metrics[0].keys())
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")

        # 写入数据
        for row_idx, metric in enumerate(metrics, start=2):
            for col_idx, key in enumerate(headers, start=1):
                ws.cell(row=row_idx, column=col_idx, value=metric[key])

        # 自动调整列宽
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except (AttributeError, ValueError, TypeError) as e:
                    # 忽略无法获取值的单元格
                    logger.debug(f"无法获取单元格值: {e}")
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column].width = adjusted_width

    def _create_statistics_sheet(self, wb, session_id: int):
        """创建统计数据工作表"""
        ws = wb.create_sheet("统计数据")

        stats = self.metrics_repo.get_statistics(session_id)
        if not stats:
            return

        # 标题
        ws['A1'] = "性能指标统计"
        ws['A1'].font = Font(bold=True, size=14)

        row = 3
        header_font = Font(bold=True, size=11)

        # 遍历各项指标统计
        for metric_name, metric_stats in stats.items():
            ws[f'A{row}'] = metric_name.upper()
            ws[f'A{row}'].font = header_font
            row += 1

            for stat_key, stat_value in metric_stats.items():
                ws[f'A{row}'] = stat_key
                ws[f'B{row}'] = stat_value
                row += 1

            row += 1

        # 设置列宽
        ws.column_dimensions['A'].width = 20
        ws.column_dimensions['B'].width = 20

    def _create_alerts_sheet(self, wb, session_id: int):
        """创建告警记录工作表"""
        ws = wb.create_sheet("告警记录")

        alerts = self.db.get_alerts(session_id=session_id)
        if not alerts:
            ws['A1'] = "无告警记录"
            return

        # 写入表头
        headers = list(alerts[0].keys())
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")

        # 写入数据
        for row_idx, alert in enumerate(alerts, start=2):
            for col_idx, key in enumerate(headers, start=1):
                ws.cell(row=row_idx, column=col_idx, value=alert[key])

        # 自动调整列宽
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except (AttributeError, ValueError, TypeError) as e:
                    # 忽略无法获取值的单元格
                    logger.debug(f"无法获取单元格值: {e}")
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column].width = adjusted_width

    def _create_trends_sheet(self, wb, session_id: int):
        """创建趋势数据工作表"""
        ws = wb.create_sheet("趋势数据")

        # 获取FPS趋势
        fps_trend = self.metrics_repo.get_fps_trend(session_id)
        ws['A1'] = "FPS趋势"
        ws['A1'].font = Font(bold=True)
        ws['B1'] = "FPS值"

        row = 2
        for timestamp, fps in fps_trend:
            ws[f'A{row}'] = timestamp.strftime('%Y-%m-%d %H:%M:%S')
            ws[f'B{row}'] = fps
            row += 1

        # 获取内存趋势
        row += 2
        ws[f'A{row}'] = "内存趋势"
        ws[f'A{row}'].font = Font(bold=True)
        ws[f'B{row}'] = "内存值(MB)"
        row += 1

        memory_trend = self.metrics_repo.get_memory_trend(session_id)
        for timestamp, memory in memory_trend:
            ws[f'A{row}'] = timestamp.strftime('%Y-%m-%d %H:%M:%S')
            ws[f'B{row}'] = memory
            row += 1

        # 设置列宽
        ws.column_dimensions['A'].width = 25
        ws.column_dimensions['B'].width = 15

    def export_batch(self, session_ids: List[int], output_dir: str,
                     format: str = 'excel') -> List[str]:
        """
        批量导出多个会话

        Args:
            session_ids: 会话ID列表
            output_dir: 输出目录
            format: 导出格式 ('csv', 'json', 'excel')

        Returns:
            成功导出的文件路径列表
        """
        success_files = []

        for session_id in session_ids:
            session = self.db.get_session(session_id)
            if not session:
                continue

            # 生成文件名
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            package_name = session['package_name'].replace('.', '_')
            device_name = session['device_id'].replace(':', '_')

            if format == 'csv':
                filename = f"{timestamp}_{device_name}_{package_name}_session{session_id}.csv"
                filepath = os.path.join(output_dir, filename)
                if self.export_to_csv(session_id, filepath):
                    success_files.append(filepath)

            elif format == 'json':
                filename = f"{timestamp}_{device_name}_{package_name}_session{session_id}.json"
                filepath = os.path.join(output_dir, filename)
                if self.export_to_json(session_id, filepath):
                    success_files.append(filepath)

            elif format == 'excel':
                filename = f"{timestamp}_{device_name}_{package_name}_session{session_id}.xlsx"
                filepath = os.path.join(output_dir, filename)
                if self.export_to_excel(session_id, filepath):
                    success_files.append(filepath)

        return success_files

    def generate_report(self, session_id: int, output_dir: str) -> Optional[str]:
        """
        生成性能测试报告（Markdown格式）

        Args:
            session_id: 会话ID
            output_dir: 输出目录

        Returns:
            报告文件路径，失败返回None
        """
        try:
            os.makedirs(output_dir, exist_ok=True)

            # 获取数据
            overview = self.session_repo.get_session_overview(session_id)
            if not overview:
                return None

            stats = self.metrics_repo.get_statistics(session_id)
            alert_summary = self.alert_repo.get_alert_summary(session_id)
            issues = self.alert_repo.get_performance_issues(session_id)

            # 生成报告文件名
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            package_name = overview['session_info']['package_name'].replace('.', '_')
            report_path = os.path.join(output_dir, f"report_{timestamp}_{package_name}.md")

            # 生成Markdown报告
            with open(report_path, 'w', encoding='utf-8') as f:
                # 标题
                f.write(f"# 性能测试报告\n\n")
                f.write(f"**生成时间**: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n\n")

                # 1. 监控概览
                f.write("## 1. 监控概览\n\n")
                session_info = overview['session_info']

                f.write(f"- **设备名称**: {session_info['device_name']}\n")
                f.write(f"- **平台**: {session_info['platform']}\n")
                f.write(f"- **应用包名**: {session_info['package_name']}\n")
                f.write(f"- **开始时间**: {session_info['start_time']}\n")
                f.write(f"- **结束时间**: {session_info.get('end_time', '进行中')}\n")
                f.write(f"- **监控时长**: {session_info['duration_formatted']}\n")
                f.write(f"- **采样间隔**: {session_info['sample_interval']}ms\n")
                f.write(f"- **样本总数**: {overview['data_summary']['total_samples']}\n\n")

                # 2. 核心指标统计
                f.write("## 2. 核心指标统计\n\n")

                if 'fps' in stats:
                    fps = stats['fps']
                    f.write("### FPS (帧率)\n\n")
                    f.write(f"- **平均值**: {fps['avg']} fps\n")
                    f.write(f"- **最大值**: {fps['max']} fps\n")
                    f.write(f"- **最小值**: {fps['min']} fps\n")
                    f.write(f"- **中位数**: {fps['median']} fps\n")
                    if fps.get('stdev'):
                        f.write(f"- **标准差**: {fps['stdev']}\n")
                    f.write("\n")

                if 'cpu_app' in stats:
                    cpu = stats['cpu_app']
                    f.write("### CPU 使用率\n\n")
                    f.write(f"- **平均值**: {cpu['avg']}%\n")
                    f.write(f"- **最大值**: {cpu['max']}%\n")
                    f.write(f"- **最小值**: {cpu['min']}%\n")
                    f.write(f"- **中位数**: {cpu['median']}%\n\n")

                if 'memory_pss' in stats:
                    memory = stats['memory_pss']
                    f.write("### 内存使用 (PSS)\n\n")
                    f.write(f"- **平均值**: {memory['avg']} MB\n")
                    f.write(f"- **最大值**: {memory['max']} MB\n")
                    f.write(f"- **最小值**: {memory['min']} MB\n")
                    f.write(f"- **中位数**: {memory['median']} MB\n\n")

                if 'network_up' in stats or 'network_down' in stats:
                    f.write("### 网络流量\n\n")
                    if 'network_up' in stats:
                        up = stats['network_up']
                        f.write(f"- **上行**: 平均 {up['avg']} KB/s, 最大 {up['max']} KB/s, 总计 {up['total_mb']} MB\n")
                    if 'network_down' in stats:
                        down = stats['network_down']
                        f.write(f"- **下行**: 平均 {down['avg']} KB/s, 最大 {down['max']} KB/s, 总计 {down['total_mb']} MB\n")
                    f.write("\n")

                # 3. 异常告警汇总
                f.write("## 3. 异常告警汇总\n\n")

                if alert_summary['total_count'] > 0:
                    f.write(f"- **告警总数**: {alert_summary['total_count']}\n")
                    f.write(f"- **未解决**: {alert_summary['unresolved_count']}\n\n")

                    f.write("### 按类型统计\n\n")
                    for alert_type, count in alert_summary['by_type'].items():
                        f.write(f"- **{alert_type}**: {count} 次\n")
                    f.write("\n")

                    f.write("### 按严重程度统计\n\n")
                    for severity, count in alert_summary['by_severity'].items():
                        f.write(f"- **{severity}**: {count} 次\n")
                    f.write("\n")

                    if alert_summary['recent_alerts']:
                        f.write("### 最近告警\n\n")
                        for alert in alert_summary['recent_alerts'][:10]:
                            f.write(f"- **{alert['type']}** [{alert['severity']}]: {alert['description']}\n")
                            f.write(f"  - 时间: {alert['timestamp']}\n\n")
                else:
                    f.write("本次监控未产生告警。\n\n")

                # 4. 性能问题分析
                f.write("## 4. 性能问题分析\n\n")

                if issues['critical_issues']:
                    f.write("### 严重问题\n\n")
                    for issue in issues['critical_issues']:
                        f.write(f"- **{issue['type']}**: {issue['description']}\n")
                        f.write(f"  - 指标: {issue['metric']}\n")
                        f.write(f"  - 当前值: {issue['current_value']}, 阈值: {issue['threshold']}\n")
                        f.write(f"  - 时间: {issue['timestamp']}\n\n")

                if issues['warnings']:
                    f.write("### 警告信息\n\n")
                    for issue in issues['warnings']:
                        f.write(f"- **{issue['type']}**: {issue['description']}\n")
                        f.write(f"  - 当前值: {issue['current_value']}, 阈值: {issue['threshold']}\n\n")

                if issues['recommendations']:
                    f.write("### 优化建议\n\n")
                    for rec in issues['recommendations']:
                        f.write(f"- **{rec['category']}**: {rec['suggestion']}\n")
                    f.write("\n")

                # 5. 总结与建议
                f.write("## 5. 总结\n\n")

                if alert_summary['total_count'] == 0:
                    f.write("本次性能测试表现良好，未发现异常情况。\n\n")
                elif alert_summary['unresolved_count'] == 0:
                    f.write(f"本次性能测试共产生 {alert_summary['total_count']} 个告警，但已全部解决。\n\n")
                else:
                    f.write(f"本次性能测试共产生 {alert_summary['total_count']} 个告警，"
                           f"其中 {alert_summary['unresolved_count']} 个尚未解决。建议优先处理严重问题。\n\n")

                f.write("---\n")
                f.write("*本报告由 Insight-Eye 性能监控工具自动生成*\n")

            print(f"测试报告生成成功: {report_path}")
            return report_path

        except Exception as e:
            print(f"生成报告失败: {e}")
            return None

    def export_to_pdf(self, session_id: int, filepath: str,
                      include_charts: bool = False,
                      chart_widgets: dict = None) -> bool:
        """
        导出为PDF格式（包含图表）

        Args:
            session_id: 会话ID
            filepath: 输出文件路径
            include_charts: 是否包含图表截图
            chart_widgets: 图表组件字典 {'fps': widget, 'cpu': widget, ...}

        Returns:
            是否成功
        """
        try:
            # 导入 reportlab
            try:
                from reportlab.lib.pagesizes import A4, letter
                from reportlab.lib.units import inch
                from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image, PageBreak
                from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
                from reportlab.lib import colors
                from reportlab.pdfbase import pdfmetrics
                from reportlab.pdfbase.ttfonts import TTFont
                from reportlab.lib.enums import TA_CENTER, TA_LEFT
            except ImportError:
                logger.error("需要安装 reportlab 库: pip install reportlab")
                return False

            # 验证文件路径
            validated_path = validate_filepath(filepath)

            # 确保目录存在
            os.makedirs(os.path.dirname(validated_path), exist_ok=True)

            # 获取会话信息
            session_overview = self.session_repo.get_session_overview(session_id)
            if not session_overview:
                logger.warning(f"会话 {session_id} 不存在")
                return False

            # 创建 PDF 文档
            doc = SimpleDocTemplate(
                validated_path,
                pagesize=A4,
                rightMargin=0.75 * inch,
                leftMargin=0.75 * inch,
                topMargin=0.75 * inch,
                bottomMargin=0.75 * inch
            )

            # 创建样式
            styles = getSampleStyleSheet()
            title_style = ParagraphStyle(
                'CustomTitle',
                parent=styles['Heading1'],
                fontSize=24,
                textColor=colors.HexColor('#00d4ff'),
                spaceAfter=30,
                alignment=TA_CENTER
            )
            heading_style = ParagraphStyle(
                'CustomHeading',
                parent=styles['Heading2'],
                fontSize=16,
                textColor=colors.HexColor('#00d4ff'),
                spaceAfter=12,
                spaceBefore=20
            )
            normal_style = styles['BodyText']

            # 构建 PDF 内容
            story = []

            # 标题
            story.append(Paragraph("Insight-Eye 性能监控报告", title_style))
            story.append(Spacer(1, 0.2 * inch))

            # 1. 会话信息
            story.append(Paragraph("1. 监控概览", heading_style))

            session_info = session_overview['session_info']
            info_data = [
                ['项目', '详情'],
                ['设备名称', session_info['device_name']],
                ['平台', session_info['platform']],
                ['应用包名', session_info['package_name']],
                ['开始时间', session_info['start_time']],
                ['结束时间', session_info.get('end_time', '进行中')],
                ['监控时长', session_info['duration_formatted']],
                ['采样间隔', f"{session_info['sample_interval']}ms"],
            ]

            info_table = Table(info_data, colWidths=[2 * inch, 4 * inch])
            info_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (1, 0), colors.HexColor('#1a1f2e')),
                ('TEXTCOLOR', (0, 0), (1, 0), colors.whitesmoke),
                ('FONTNAME', (0, 0), (1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (1, 0), 12),
                ('BOTTOMPADDING', (0, 0), (1, 0), 12),
                ('BACKGROUND', (0, 1), (-1, -1), colors.HexColor('#141414')),
                ('TEXTCOLOR', (0, 1), (-1, -1), colors.HexColor('#e0e6ed')),
                ('GRID', (0, 0), (-1, -1), 1, colors.HexColor('#3b4252')),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor('#141414'), colors.HexColor('#1a1f2e')]),
            ]))
            story.append(info_table)
            story.append(Spacer(1, 0.3 * inch))

            # 2. 性能指标统计
            story.append(Paragraph("2. 性能指标统计", heading_style))

            stats = self.metrics_repo.get_statistics(session_id)

            # FPS 统计
            if 'fps' in stats:
                fps = stats['fps']
                fps_data = [
                    ['指标', '数值'],
                    ['平均 FPS', f"{fps['avg']:.1f}"],
                    ['最大 FPS', f"{fps['max']:.1f}"],
                    ['最小 FPS', f"{fps['min']:.1f}"],
                    ['中位数', f"{fps['median']:.1f}"],
                ]
                fps_table = Table(fps_data, colWidths=[2.5 * inch, 3.5 * inch])
                fps_table.setStyle(TableStyle([
                    ('BACKGROUND', (0, 0), (1, 0), colors.HexColor('#ffb400')),
                    ('TEXTCOLOR', (0, 0), (1, 0), colors.black),
                    ('FONTNAME', (0, 0), (1, 0), 'Helvetica-Bold'),
                    ('BACKGROUND', (0, 1), (-1, -1), colors.HexColor('#141414')),
                    ('TEXTCOLOR', (0, 1), (-1, -1), colors.HexColor('#e0e6ed')),
                    ('GRID', (0, 0), (-1, -1), 1, colors.HexColor('#3b4252')),
                ]))
                story.append(Paragraph("FPS (帧率)", ParagraphStyle('SubHeading', parent=styles['Heading3'], fontSize=14)))
                story.append(fps_table)
                story.append(Spacer(1, 0.2 * inch))

            # CPU 统计
            if 'cpu_app' in stats:
                cpu = stats['cpu_app']
                cpu_data = [
                    ['指标', '数值'],
                    ['平均 CPU', f"{cpu['avg']:.1f}%"],
                    ['最大 CPU', f"{cpu['max']:.1f}%"],
                    ['最小 CPU', f"{cpu['min']:.1f}%"],
                    ['中位数', f"{cpu['median']:.1f}%"],
                ]
                cpu_table = Table(cpu_data, colWidths=[2.5 * inch, 3.5 * inch])
                cpu_table.setStyle(TableStyle([
                    ('BACKGROUND', (0, 0), (1, 0), colors.HexColor('#00f2ff')),
                    ('TEXTCOLOR', (0, 0), (1, 0), colors.black),
                    ('FONTNAME', (0, 0), (1, 0), 'Helvetica-Bold'),
                    ('BACKGROUND', (0, 1), (-1, -1), colors.HexColor('#141414')),
                    ('TEXTCOLOR', (0, 1), (-1, -1), colors.HexColor('#e0e6ed')),
                    ('GRID', (0, 0), (-1, -1), 1, colors.HexColor('#3b4252')),
                ]))
                story.append(Paragraph("CPU 使用率", ParagraphStyle('SubHeading', parent=styles['Heading3'], fontSize=14)))
                story.append(cpu_table)
                story.append(Spacer(1, 0.2 * inch))

            # 内存统计
            if 'memory_pss' in stats:
                memory = stats['memory_pss']
                memory_data = [
                    ['指标', '数值'],
                    ['平均内存', f"{memory['avg']:.0f} MB"],
                    ['最大内存', f"{memory['max']:.0f} MB"],
                    ['最小内存', f"{memory['min']:.0f} MB"],
                    ['中位数', f"{memory['median']:.0f} MB"],
                ]
                memory_table = Table(memory_data, colWidths=[2.5 * inch, 3.5 * inch])
                memory_table.setStyle(TableStyle([
                    ('BACKGROUND', (0, 0), (1, 0), colors.HexColor('#7000ff')),
                    ('TEXTCOLOR', (0, 0), (1, 0), colors.white),
                    ('FONTNAME', (0, 0), (1, 0), 'Helvetica-Bold'),
                    ('BACKGROUND', (0, 1), (-1, -1), colors.HexColor('#141414')),
                    ('TEXTCOLOR', (0, 1), (-1, -1), colors.HexColor('#e0e6ed')),
                    ('GRID', (0, 0), (-1, -1), 1, colors.HexColor('#3b4252')),
                ]))
                story.append(Paragraph("内存使用 (PSS)", ParagraphStyle('SubHeading', parent=styles['Heading3'], fontSize=14)))
                story.append(memory_table)
                story.append(Spacer(1, 0.2 * inch))

            # 网络统计
            if 'network_down' in stats:
                down = stats['network_down']
                network_data = [
                    ['指标', '数值'],
                    ['平均下行速率', f"{down['avg']:.1f} KB/s"],
                    ['最大下行速率', f"{down['max']:.1f} KB/s"],
                    ['总下载量', f"{down.get('total_mb', 0):.1f} MB"],
                ]
                network_table = Table(network_data, colWidths=[2.5 * inch, 3.5 * inch])
                network_table.setStyle(TableStyle([
                    ('BACKGROUND', (0, 0), (1, 0), colors.HexColor('#00ff87')),
                    ('TEXTCOLOR', (0, 0), (1, 0), colors.black),
                    ('FONTNAME', (0, 0), (1, 0), 'Helvetica-Bold'),
                    ('BACKGROUND', (0, 1), (-1, -1), colors.HexColor('#141414')),
                    ('TEXTCOLOR', (0, 1), (-1, -1), colors.HexColor('#e0e6ed')),
                    ('GRID', (0, 0), (-1, -1), 1, colors.HexColor('#3b4252')),
                ]))
                story.append(Paragraph("网络流量", ParagraphStyle('SubHeading', parent=styles['Heading3'], fontSize=14)))
                story.append(network_table)
                story.append(Spacer(1, 0.2 * inch))

            # 3. 告警摘要
            alert_summary = self.alert_repo.get_alert_summary(session_id)
            if alert_summary['total_count'] > 0:
                story.append(Paragraph("3. 异常告警汇总", heading_style))

                alert_data = [
                    ['类型', '次数'],
                ]
                for alert_type, count in alert_summary['by_type'].items():
                    alert_data.append([alert_type, str(count)])

                alert_table = Table(alert_data, colWidths=[3 * inch, 3 * inch])
                alert_table.setStyle(TableStyle([
                    ('BACKGROUND', (0, 0), (1, 0), colors.HexColor('#ef4444')),
                    ('TEXTCOLOR', (0, 0), (1, 0), colors.white),
                    ('FONTNAME', (0, 0), (1, 0), 'Helvetica-Bold'),
                    ('BACKGROUND', (0, 1), (-1, -1), colors.HexColor('#141414')),
                    ('TEXTCOLOR', (0, 1), (-1, -1), colors.HexColor('#e0e6ed')),
                    ('GRID', (0, 0), (-1, -1), 1, colors.HexColor('#3b4252')),
                ]))
                story.append(alert_table)
                story.append(Spacer(1, 0.2 * inch))

            # 4. 图表（可选）
            if include_charts and chart_widgets:
                story.append(PageBreak())
                story.append(Paragraph("4. 性能趋势图表", heading_style))

                # 捕获图表并添加到 PDF
                for chart_name, chart_widget in chart_widgets.items():
                    if chart_widget is not None:
                        try:
                            # 捕获图表为图片
                            pixmap = chart_widget.grab()
                            from PyQt6.QtCore import QBuffer
                            from PyQt6.QtGui import QImage

                            # 保存为临时字节流
                            buffer = QBuffer()
                            buffer.open(QBuffer.OpenModeFlag.ReadWrite)
                            pixmap.save(buffer, 'PNG')
                            image_data = buffer.data()
                            buffer.close()

                            # 创建 reportlab Image
                            img = Image(image_data, width=6 * inch, height=3 * inch)
                            story.append(Paragraph(f"{chart_name.upper()} 趋势图", ParagraphStyle('ChartTitle', parent=styles['Heading4'], fontSize=12)))
                            story.append(img)
                            story.append(Spacer(1, 0.2 * inch))

                        except Exception as e:
                            logger.warning(f"无法捕获图表 {chart_name}: {e}")

            # 生成 PDF
            doc.build(story)
            logger.info(f"PDF导出成功: {validated_path}")
            return True

        except Exception as e:
            logger.error(f"PDF导出失败: {e}")
            import traceback
            traceback.print_exc()
            return False
